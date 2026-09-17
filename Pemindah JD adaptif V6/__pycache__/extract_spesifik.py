
import os
import glob
import textwrap
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

def extract_specific_tasks(source_sheet):
    """
    Mengambil poin Tugas Pokok Spesifik dari kolom V-AO (22-41)
    tanpa tergantung warna sel (bekerja pada sel 'no fill' maupun berwarna).
    """
    items = []
    current_item = None
    is_in_spesifik = False
    
    start_col = 22  # Kolom V
    end_col = 41    # Kolom AO

    for row in range(1, source_sheet.max_row + 1):
        row_vals = []
        for col in range(start_col, end_col + 1):
            cell = source_sheet.cell(row=row, column=col)
            if cell.value is not None:
                v = str(cell.value).strip()
                if v and v not in row_vals:
                    row_vals.append(v)
                    
        line_str = " ".join(row_vals).strip()
        if not line_str:
            continue

        # 1. DETEKSI MULAI (HEADER SPESIFIK)
        if not is_in_spesifik:
            if "SPESIFIK" in line_str.upper() and len(line_str) < 30:
                is_in_spesifik = True
                continue

        # 2. PROSES AREA TUGAS SPESIFIK
        if is_in_spesifik:
            stop_keywords = [
                "PENDIDIKAN", "WEWENANG", "HUBUNGAN KERJA", 
                "TANGGUNG JAWAB", "KETERKAITAN DENGAN PIHAK LAIN",
                "TANGGUNGJAWAB", "PERSYARATAN"
            ]
            if any(stop_kw in line_str.upper() for stop_kw in stop_keywords):
                break
                
            first_word = row_vals[0].replace('.', '').strip()
            
            if first_word.isdigit():
                col_no = int(first_word)
                text_combined = " ".join(row_vals[1:]).strip()
                
                if current_item:
                    items.append(current_item)
                current_item = {"no": col_no, "text": text_combined}
            elif current_item:
                current_item["text"] += " " + line_str

    if current_item:
        items.append(current_item)
        
    return items


def insert_rows_and_preserve_layout(ws, row_idx, amount):
    """
    Menyisipkan baris penuh di row_idx, menggeser merged cells,
    DAN menggeser tinggi baris (row heights) agar tabel bawah tidak gepeng.
    """
    if amount <= 0:
        return

    # 1. Simpan tinggi baris (row height) lama dari row_idx ke bawah
    old_heights = {}
    for r in range(row_idx, ws.max_row + 1):
        if r in ws.row_dimensions and ws.row_dimensions[r].height is not None:
            old_heights[r] = ws.row_dimensions[r].height

    # 2. Simpan & unmerge temporary merged cells dari row_idx ke bawah
    merged_ranges = list(ws.merged_cells.ranges)
    ranges_to_shift = []
    for rng in merged_ranges:
        if rng.min_row >= row_idx:
            ranges_to_shift.append(rng)
            ws.unmerge_cells(range_string=str(rng))

    # 3. Sisipkan entire row
    ws.insert_rows(row_idx, amount)

    # 4. Merge kembali dengan koordinat baru yang sudah digeser (+ amount)
    for rng in ranges_to_shift:
        ws.merge_cells(
            start_row=rng.min_row + amount,
            start_column=rng.min_col,
            end_row=rng.max_row + amount,
            end_column=rng.max_col
        )

    # 5. TERAPKAN KEMBALI TINGGI BARIS (ROW HEIGHT) KE POSISI BARU
    for r in range(row_idx, row_idx + amount):
        ws.row_dimensions[r].height = 16.5

    for old_r, h in old_heights.items():
        ws.row_dimensions[old_r + amount].height = h


def process_single_file(source_path, template_path, output_path):
    try:
        wb_src = openpyxl.load_workbook(source_path, data_only=True)
        sheet_src = wb_src.active
        
        wb_tpl = openpyxl.load_workbook(template_path)
        sheet_tpl = wb_tpl.active
        
        data_spesifik = extract_specific_tasks(sheet_src)
        if not data_spesifik:
            print(f"⚠️ Warning: Tidak ada poin spesifik terdeteksi di {os.path.basename(source_path)}")
            return
            
        start_row_tgt = 38
        max_default_rows = 9  # Row 38 sampai 46
        
        # PATOKAN KARAKTER DIPASANG 110 AGAR TIDAK MENABRAK BATAS MURNI KOLOM AO
        CHARS_PER_LINE = 110

        # 1. HITUNG JUMLAH BARIS REAL DENGAN PEMOTONGAN KATA PAKSA (PAKAI BREAK_LONG_WORDS=TRUE)
        justified_rows = []
        for idx, item in enumerate(data_spesifik, start=1):
            full_text = item["text"]
            
            # PAKSA POTONG WALAUPUN TEKS TIDAK PUNYA SPASI!
            lines = textwrap.wrap(
                full_text, 
                width=CHARS_PER_LINE, 
                break_long_words=True, 
                break_on_hyphens=False
            ) or [""]

            for l_idx, line in enumerate(lines):
                justified_rows.append({
                    "no": f"{idx}." if l_idx == 0 else "",  # Nomor cuma di baris pertama
                    "text": line
                })

        total_needed_rows = len(justified_rows)

        # 2. INSERT ENTIRE ROW JIKA BARIS YANG DIBUTUHKAN > 9 BARIS
        extra_rows = 0
        if total_needed_rows > max_default_rows:
            extra_rows = total_needed_rows - max_default_rows
            insert_rows_and_preserve_layout(sheet_tpl, row_idx=46, amount=extra_rows)

        font_arial_10 = Font(name="Arial", size=10, bold=False)
        double_side = Side(border_style="double", color="000000")
        last_item_row = start_row_tgt + total_needed_rows - 1

        # 3. PASTE HASIL JUSTIFY PAKSA KE SEL
        for idx, row_data in enumerate(justified_rows, start=1):
            curr_row = start_row_tgt + idx - 1
            sheet_tpl.row_dimensions[curr_row].height = 16.5

            # A. Nomor di Kolom B
            cell_no = sheet_tpl.cell(row=curr_row, column=2)
            cell_no.value = row_data["no"]
            cell_no.font = font_arial_10
            cell_no.alignment = Alignment(horizontal='center', vertical='center', wrap_text=False)
            
            # B. Teks di Kolom C (Dijamin TIDAK BISA TEMBUS lagi karena dipotong paksa di 110 karakter)
            cell_text = sheet_tpl.cell(row=curr_row, column=3)
            cell_text.value = row_data["text"]
            cell_text.font = font_arial_10
            cell_text.alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)

            # C. Tambal Garis Vertikal Ganda (Double Border)
            sheet_tpl.cell(row=curr_row, column=2).border = Border(left=double_side)
            sheet_tpl.cell(row=curr_row, column=41).border = Border(right=double_side)
            sheet_tpl.cell(row=curr_row, column=42).border = Border(left=double_side)
            sheet_tpl.cell(row=curr_row, column=63).border = Border(right=double_side)

        # 4. HILANGKAN BOTTOM BORDER HANYA DI BARIS TERAKHIR POIN
        for c in range(2, 64):
            cell = sheet_tpl.cell(row=last_item_row, column=c)
            curr_b = cell.border
            sheet_tpl.cell(row=last_item_row, column=c).border = Border(
                left=curr_b.left if curr_b else None,
                right=curr_b.right if curr_b else None,
                top=curr_b.top if curr_b else None,
                bottom=None
            )

        wb_tpl.save(output_path)
        print(f"✅ [BERHASIL] {os.path.basename(source_path)} -> {os.path.basename(output_path)} ({total_needed_rows} baris terisi aman!)")

    except Exception as e:
        print(f"❌ [GAGAL] Error memproses {os.path.basename(source_path)}: {str(e)}")


def batch_process():
    input_dir = "input_jobdesc_lama"
    template_path = "Template_JD_baru_Gol_X.xlsx"
    output_dir = "output_jobdesc_baru"
    
    os.makedirs(output_dir, exist_ok=True)
    excel_files = glob.glob(os.path.join(input_dir, "*.xlsx"))
    
    for file_path in excel_files:
        if os.path.basename(file_path).startswith("~$"):
            continue
        out_name = f"Baru_{os.path.basename(file_path)}"
        process_single_file(file_path, template_path, os.path.join(output_dir, out_name))


if __name__ == "__main__":
    batch_process()

    
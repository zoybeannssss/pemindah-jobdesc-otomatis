import os
import openpyxl
from openpyxl.styles import PatternFill

def extract_tugas_spesifik(source_sheet):
    """Mengekstrak list poin Tugas Pokok Spesifik dari Job Desc Lama (Kolom V & W)"""
    items = []
    current_item = None
    is_spesifik = False
    
    for row in range(1, source_sheet.max_row + 1):
        v_val = source_sheet.cell(row=row, column=22).value  # Kolom V
        w_val = source_sheet.cell(row=row, column=23).value  # Kolom W
        
        v_str = str(v_val).strip() if v_val is not None else ""
        w_str = str(w_val).strip() if w_val is not None else ""
        
        # Pemicu awal pembacaan
        if "Spesifik :" in v_str:
            is_spesifik = True
            continue
            
        # Pemicu akhir pembacaan
        if is_spesifik and ("JOB SPECIFICATION" in v_str or "WEWENANG" in str(source_sheet.cell(row=row, column=47).value or "")):
            break
            
        if is_spesifik:
            clean_v = v_str.replace('.', '').strip()
            if clean_v.isdigit():
                if current_item:
                    items.append(current_item)
                current_item = {"no": clean_v, "text": w_str}
            elif w_str and current_item:
                current_item["text"] += " " + w_str
                
    if current_item:
        items.append(current_item)
    return items


def execute_migration(source_path, template_path, output_path):
    print("🚀 Memulai proses pemindahan Job Desc...")
    
    # 1. Load File Excel
    wb_src = openpyxl.load_workbook(source_path, data_only=True)
    sheet_src = wb_src.active
    
    wb_tpl = openpyxl.load_workbook(template_path)
    sheet_tpl = wb_tpl.active
    
    # 2. Extract Data Spesifik
    data_spesifik = extract_tugas_spesifik(sheet_src)
    print(f"✅ Extracted {len(data_spesifik)} poin dari Job Desc Lama.")
    
    # 3. Cari Variabel "TugasPokokSpecificMoved"
    target_row = None
    target_col = None
    
    for r in range(1, sheet_tpl.max_row + 1):
        for c in range(1, sheet_tpl.max_column + 1):
            val = sheet_tpl.cell(row=r, column=c).value
            if val and "TugasPokokSpecificMoved" in str(val):
                target_row = r
                target_col = c
                break
        if target_row:
            break
            
    if not target_row:
        print("❌ Error: Variable 'TugasPokokSpecificMoved' tidak ditemukan di template!")
        return

    print(f"📍 Target variabel ditemukan pada Sel Row {target_row}, Column {target_col}")
    
    # 4. Sisipkan Baris Dinamis jika Jumlah Poin > 1
    num_items = len(data_spesifik)
    if num_items > 1:
        sheet_tpl.insert_rows(target_row + 1, amount=num_items - 1)
        
    # 5. Inject Data & Hapus Background Warna (Termasuk Kuning)
    fill_clear = PatternFill(fill_type=None)
    
    for idx, item in enumerate(data_spesifik):
        curr_row = target_row + idx
        
        # Kolom Nomor (B)
        cell_no = sheet_tpl.cell(row=curr_row, column=target_col)
        cell_no.value = f"{item['no']}."
        cell_no.fill = fill_clear
        
        # Kolom Deskripsi (C)
        cell_text = sheet_tpl.cell(row=curr_row, column=target_col + 1)
        cell_text.value = item['text']
        cell_text.fill = fill_clear
        
    # 6. Simpan File Output
    wb_tpl.save(output_path)
    print(f"🎉 Sukses! File hasil migrasi berhasil disimpan di: {output_path}")

if __name__ == "__main__":
    SRC_FILE = "files/Job desc Lama.xlsx"
    TPL_FILE = "files/Template Job Desc Baru - Output.xlsx"
    OUT_FILE = "files/Hasil_Migrasi_JobDesc.xlsx"
    
    execute_migration(SRC_FILE, TPL_FILE, OUT_FILE)
